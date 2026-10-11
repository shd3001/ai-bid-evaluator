"""Local parsers preserve source locations; quality failures are structured.

pdfplumber is selected over the generic MarkItDown PDF conversion because
it exposes pages/tables. OCR is optional and must return text AND confidence.
No OCR engine is silently downloaded or claimed to be available.
"""
import csv
import hashlib
import io
import json
import math
from pathlib import Path
from typing import Callable
from zipfile import ZipFile
from xml.etree import ElementTree
from ai_bid_evaluator.domain.documents import ParsedDocument, DocumentChunk
from ai_bid_evaluator.domain.exceptions import UnsupportedFormat

SUPPORTED_FORMATS = {".pdf", ".docx", ".xlsx", ".xlsm", ".csv", ".txt", ".md", ".json", ".html", ".htm"}
OCRBackend = Callable[[Path, int], tuple[str, float]]


def ingest_document(path: Path, document_id: str, *, ocr: OCRBackend | None = None) -> ParsedDocument:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_FORMATS:
        raise UnsupportedFormat("格式没有已验证解析器，请转换为 PDF/DOCX/XLSX 或文本")
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    chunks, diagnostics = [], []
    status, parser = "ready", suffix.lstrip(".")

    def add(text, **location):
        if text and text.strip():
            chunks.append(DocumentChunk(document_id=document_id, chunk_id=f"{document_id}_{len(chunks)+1}",
                                        document_hash=digest, text=text, **location))

    try:
        if suffix == ".pdf":
            import pdfplumber
            parser = "pdfplumber"
            with pdfplumber.open(path) as pdf:
                for number, page in enumerate(pdf.pages, 1):
                    text = page.extract_text() or ""
                    quality = "good"
                    if not text.strip():
                        quality, status = "needs_review", "needs_review"
                        if ocr is None:
                            diagnostics.append(f"page {number}: OCR 未配置，扫描页或空页需人工检查")
                        else:
                            try:
                                text, confidence = ocr(path, number)
                                if not math.isfinite(confidence) or not 0 <= confidence <= 1:
                                    raise ValueError("invalid OCR confidence")
                                quality = "good" if confidence >= 0.85 and text.strip() else "needs_review"
                                diagnostics.append(f"page {number}: OCR confidence {confidence:.2f}")
                            except Exception:
                                text = ""
                                diagnostics.append(f"page {number}: OCR 失败")
                    add(text, page=number, quality=quality)
                    for table_number, table in enumerate(page.extract_tables(), 1):
                        for row_number, row in enumerate(table, 1):
                            add(" | ".join(str(c or "") for c in row), page=number,
                                sheet=f"table_{table_number}", row_start=row_number, row_end=row_number,
                                column_start=1, column_end=len(row), quality=quality)
        elif suffix in {".xlsx", ".xlsm"}:
            from openpyxl import load_workbook
            parser = "openpyxl"
            book = load_workbook(path, read_only=True, data_only=False)
            try:
                for sheet in book:
                    for number, row in enumerate(sheet.iter_rows(), 1):
                        values = []
                        quality = "good"
                        for cell in row:
                            value = cell.value
                            if cell.data_type == "f":
                                quality, status = "needs_review", "needs_review"
                                diagnostics.append(f"sheet {sheet.title} row {number}: 公式需核实计算值")
                            values.append(str(value) if value is not None else "")
                        if any(values):
                            add(" | ".join(values), sheet=sheet.title, row_start=number, row_end=number,
                                column_start=1, column_end=len(row), quality=quality)
            finally:
                book.close()
        elif suffix == ".docx":
            parser = "docx_xml"
            ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
            with ZipFile(path) as archive:
                root = ElementTree.fromstring(archive.read("word/document.xml"))
            body = root.find("w:body", ns)
            for index, node in enumerate(body, 1):
                if node.tag.endswith("}tbl"):
                    for number, row in enumerate(node.findall("w:tr", ns), 1):
                        cells = ["".join(c.itertext()) for c in row.findall("w:tc", ns)]
                        add(" | ".join(cells), sheet=f"table_{index}", row_start=number, row_end=number,
                            column_start=1, column_end=len(cells))
                else:
                    add("".join(n.text or "" for n in node.findall(".//w:t", ns)))
            diagnostics.append("DOCX 块和表格可定位；未伪造排版页码")
        elif suffix == ".csv":
            for number, row in enumerate(csv.reader(io.StringIO(path.read_text(encoding="utf-8-sig"))), 1):
                add(" | ".join(row), sheet="csv", row_start=number, row_end=number, column_start=1, column_end=len(row))
        elif suffix in {".html", ".htm"}:
            from bs4 import BeautifulSoup
            soup = BeautifulSoup(path.read_text(encoding="utf-8"), "html.parser")
            for node in soup(["script", "style", "nav", "footer"]):
                node.decompose()
            add(soup.get_text("\n", strip=True))
        else:
            add(path.read_text(encoding="utf-8-sig"))
    except Exception as exc:
        # Do not turn error messages into input or leak absolute paths.
        status = "failed"
        diagnostics.append(f"解析失败 ({type(exc).__name__})")
    if not chunks and status == "ready":
        status = "failed"
        diagnostics.append("正文为空")
    if any("\ufffd" in c.text or len(c.text.strip()) < 3 for c in chunks) and status == "ready":
        status = "needs_review"
        diagnostics.append("文本质量不足")
    return ParsedDocument(document_id=document_id, document_hash=digest, format=suffix,
                          status=status, chunks=chunks, diagnostics=diagnostics, parser=parser)


def chunk_document(doc: ParsedDocument, max_tokens: int) -> list[DocumentChunk]:
    """UTF-8 byte bound is conservative for byte-level tokenizers; never truncate.

    Caller reserves prompt/output overhead separately. Locations remain attached
    to each split. No claimed exact model token count.
    """
    if max_tokens < 4:
        raise ValueError("max_tokens 至少为 4")
    chunks = []
    for source in doc.chunks:
        part, size = [], 0
        for char in source.text:
            cost = len(char.encode("utf-8"))
            if size + cost > max_tokens:
                chunks.append(source.model_copy(update={"chunk_id": f"{doc.document_id}_chunk_{len(chunks)+1}", "text": "".join(part)}))
                part, size = [], 0
            part.append(char)
            size += cost
        if part:
            chunks.append(source.model_copy(update={"chunk_id": f"{doc.document_id}_chunk_{len(chunks)+1}", "text": "".join(part)}))
    return identify_chunks(chunks)


def identify_chunks(chunks):
    # A stored evidence ID must not name different text when another run changes
    # its model budget. Include content, position and document identity.
    return [c.model_copy(update={"chunk_id": "chunk_" + hashlib.sha256(json.dumps(
        {"chunk": c.model_dump(exclude={"chunk_id"}), "index": i}, ensure_ascii=False, sort_keys=True).encode()).hexdigest()[:48]})
        for i, c in enumerate(chunks)]
